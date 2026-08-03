"""Extract the measurable house style from a reference deck.

LibreOffice cannot render in this container, so instead of looking at slides we
measure them: how many words carry a slide, how often a visual does the work,
which fonts and colours recur, and what the actual layout archetypes are.  Run
against a known-good deck to get numbers to build the next one against.

    python3 analyze_reference.py reference.pptx
"""

import collections
import re
import sys

from pptx import Presentation
from pptx.util import Length

EMU = 914400.0


def in_(v):
    return (v or 0) / EMU


def shape_kind(sh):
    """Coarse bucket: what job is this shape doing on the slide?"""
    t = str(sh.shape_type)
    if sh.has_chart if hasattr(sh, "has_chart") else False:
        return "chart"
    if "PICTURE" in t:
        return "picture"
    if "TABLE" in t or getattr(sh, "has_table", False):
        return "table"
    if "GROUP" in t:
        return "group"
    if sh.has_text_frame and sh.text_frame.text.strip():
        return "text"
    return "shape"


def run_colors(sh):
    out = []
    if not sh.has_text_frame:
        return out
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            try:
                if r.font.color and r.font.color.type is not None:
                    out.append(str(r.font.color.rgb))
            except Exception:
                pass
    return out


def fill_color(sh):
    try:
        f = sh.fill
        if f.type is not None and f.type == 1:  # solid
            return str(f.fore_color.rgb)
    except Exception:
        pass
    return None


def main(path):
    prs = Presentation(path)
    sw, sh_ = in_(prs.slide_width), in_(prs.slide_height)

    words_per, kinds_per = [], []
    fonts = collections.Counter()
    sizes = collections.Counter()
    colors = collections.Counter()
    fills = collections.Counter()
    kind_totals = collections.Counter()
    bullet_slides = 0
    visual_slides = 0
    notes_words = []
    per_slide = []

    for i, slide in enumerate(prs.slides, start=1):
        words, kinds = 0, collections.Counter()
        has_bullets = False
        for shp in slide.shapes:
            k = shape_kind(shp)
            kinds[k] += 1
            kind_totals[k] += 1
            if shp.has_text_frame:
                txt = shp.text_frame.text
                words += len(re.findall(r"\S+", txt))
                for p in shp.text_frame.paragraphs:
                    if p.level and p.level > 0:
                        has_bullets = True
                    # a buChar/buAutoNum in the paragraph properties = a real bullet
                    pPr = p._pPr
                    if pPr is not None and (pPr.find(
                            "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None
                            or pPr.find(
                            "{http://schemas.openxmlformats.org/drawingml/2006/main}buAutoNum") is not None):
                        has_bullets = True
                    for r in p.runs:
                        if r.font.name:
                            fonts[r.font.name] += 1
                        if r.font.size:
                            sizes[round(r.font.size.pt, 1)] += 1
                for c in run_colors(shp):
                    colors[c] += 1
            fc = fill_color(shp)
            if fc:
                fills[fc] += 1

        words_per.append(words)
        kinds_per.append(kinds)
        if has_bullets:
            bullet_slides += 1
        if kinds["chart"] or kinds["picture"] or kinds["table"]:
            visual_slides += 1
        if slide.has_notes_slide:
            notes_words.append(len(re.findall(
                r"\S+", slide.notes_slide.notes_text_frame.text)))
        per_slide.append((i, words, dict(kinds)))

    n = len(prs.slides)
    print(f"=== {path}")
    print(f"canvas            {sw:.2f} x {sh_:.2f} in")
    print(f"slides            {n}")
    print(f"words/slide       median {sorted(words_per)[n // 2]}, "
          f"mean {sum(words_per) / n:.0f}, max {max(words_per)}")
    print(f"slides >80 words  {sum(1 for w in words_per if w > 80)} "
          f"({100 * sum(1 for w in words_per if w > 80) / n:.0f}%)")
    print(f"slides <25 words  {sum(1 for w in words_per if w < 25)} "
          f"({100 * sum(1 for w in words_per if w < 25) / n:.0f}%)")
    print(f"with a visual     {visual_slides} ({100 * visual_slides / n:.0f}%)")
    print(f"with real bullets {bullet_slides} ({100 * bullet_slides / n:.0f}%)")
    if notes_words:
        print(f"speaker notes     {len(notes_words)}/{n} slides, "
              f"median {sorted(notes_words)[len(notes_words) // 2]} words")
    print(f"\nshape mix         {dict(kind_totals)}")
    print(f"pictures/slide    {kind_totals['picture'] / n:.2f}")
    print(f"charts/slide      {kind_totals['chart'] / n:.2f}")
    print(f"\nfonts             {fonts.most_common(6)}")
    print(f"font sizes        {sizes.most_common(12)}")
    print(f"text colors       {colors.most_common(8)}")
    print(f"fill colors       {fills.most_common(8)}")

    print("\n--- heaviest slides (words, shape mix) ---")
    for i, w, k in sorted(per_slide, key=lambda r: -r[1])[:12]:
        print(f"  s{i:<3} {w:>4}w  {k}")
    print("\n--- lightest slides ---")
    for i, w, k in sorted(per_slide, key=lambda r: r[1])[:8]:
        print(f"  s{i:<3} {w:>4}w  {k}")


if __name__ == "__main__":
    for p in sys.argv[1:] or ["reference.pptx"]:
        main(p)
        print()
