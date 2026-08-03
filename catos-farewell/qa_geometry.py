"""Geometry QA for the CATOS farewell deck.

LibreOffice cannot render in this sandbox, so instead of eyeballing images we
check the two defects that actually bite: text that will not fit its box, and
shapes that overlap or run off the canvas.  Text height is estimated from
per-font average advance widths, which is coarse but catches real overflow.
"""

import math
import sys
from pptx import Presentation
from pptx.util import Length

SLIDE_W = 13.3333
SLIDE_H = 7.5
MARGIN = 0.45          # anything closer to an edge than this is flagged
EMU = 914400.0

# Average advance width as a fraction of point size, measured for the two
# fonts this deck uses.  Bold runs a little wider.
AVG_W = {("Calibri", False): 0.465, ("Calibri", True): 0.487,
         ("Cambria", False): 0.500, ("Cambria", True): 0.523}
LINE_H = 1.22          # line box as a multiple of point size
BOX_PAD = 0.10         # pptxgenjs default internal padding, inches, per side


def in_(v):
    return (v or 0) / EMU


def est_height(tf, w_in):
    """Estimated rendered height of a text frame, in inches."""
    total = 0.0
    for para in tf.paragraphs:
        runs = [r for r in para.runs if r.text]
        if not runs:
            total += 0.10
            continue
        size = max((r.font.size.pt if r.font.size else 18) for r in runs)
        text = "".join(r.text for r in runs)
        bold = any(bool(r.font.bold) for r in runs)
        face = runs[0].font.name or "Calibri"
        key = (face if face in ("Calibri", "Cambria") else "Calibri", bold)
        cw = AVG_W[key] * size / 72.0
        usable = max(w_in - 2 * BOX_PAD, 0.2)
        per_line = max(int(usable / cw), 1)
        # explicit newlines force their own lines
        lines = sum(max(math.ceil(len(seg) / per_line), 1) for seg in text.split("\n"))
        space_after = para.space_after.pt if para.space_after else 0
        ls = para.line_spacing
        if ls is None:
            line_pt = size * LINE_H
        elif isinstance(ls, Length):      # exact spacing, stored in EMU
            line_pt = ls.pt
        else:                             # a multiple of single spacing
            line_pt = size * ls
        total += (lines * line_pt + space_after) / 72.0
    return total


def rects_overlap(a, b, tol=0.02):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    return (ax < bx + bw - tol and bx < ax + aw - tol
            and ay < by + bh - tol and by < ay + ah - tol)


def main(path):
    prs = Presentation(path)
    problems = []

    for idx, slide in enumerate(prs.slides, start=1):
        boxes = []
        for sh in slide.shapes:
            x, y = in_(sh.left), in_(sh.top)
            w, h = in_(sh.width), in_(sh.height)
            name = sh.shape_type
            label = ""
            is_footer = False
            if sh.has_text_frame and sh.text_frame.text.strip():
                label = sh.text_frame.text.strip().replace("\n", " ")[:44]
                sizes = [r.font.size.pt for p in sh.text_frame.paragraphs
                         for r in p.runs if r.font.size]
                is_footer = bool(sizes) and max(sizes) <= 10

            # off-canvas / margin
            if x < -0.01 or y < -0.01 or x + w > SLIDE_W + 0.01 or y + h > SLIDE_H + 0.01:
                problems.append(f"s{idx}: OFF-CANVAS {name} at ({x:.2f},{y:.2f}) "
                                f"{w:.2f}x{h:.2f} '{label}'")
            elif label and not is_footer and (x < MARGIN - 0.01 or y < MARGIN - 0.01
                            or x + w > SLIDE_W - MARGIN + 0.01
                            or y + h > SLIDE_H - MARGIN + 0.01):
                problems.append(f"s{idx}: TIGHT-MARGIN text at ({x:.2f},{y:.2f}) "
                                f"{w:.2f}x{h:.2f} '{label}'")

            # overflow
            if sh.has_text_frame and sh.text_frame.text.strip():
                need = est_height(sh.text_frame, w)
                if need > h + 0.06:
                    problems.append(f"s{idx}: OVERFLOW ~{need:.2f}in into {h:.2f}in box "
                                    f"at ({x:.2f},{y:.2f}) w={w:.2f} '{label}'")
                boxes.append(((x, y, w, h), label))

        # text-on-text overlap (cards are drawn first and are not text frames)
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                if rects_overlap(boxes[i][0], boxes[j][0], tol=0.05):
                    problems.append(f"s{idx}: TEXT-OVERLAP '{boxes[i][1]}' x '{boxes[j][1]}'")

    print(f"{len(prs.slides)} slides checked")
    if not problems:
        print("No geometry problems found.")
    for p in problems:
        print(p)
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1
                  else "CATOS_Farewell_Ferrara_Aug2026.pptx"))
